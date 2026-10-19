# -*- coding: utf-8 -*-
"""Build Review2_ResearchPaper.docx from paper_content.py.

Single-column A4, Times New Roman.  Run:  python3 build_paper.py
"""

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import paper_content as C

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Review2_ResearchPaper.docx")

FONT = "Times New Roman"
BODY_PT = 11.0
SMALL_PT = 10.0     # abstract, index terms, references
CAPTION_PT = 9.0    # figure captions, table text
MEASURE = 6.27      # A4 width 8.27in less 1in margins each side

SUB = re.compile(r"_\[([A-Za-z]+)\]")


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #

def style_run(r, size=BODY_PT, bold=False, italic=False):
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    # east-asian font name, else Word may substitute for non-ASCII glyphs
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    return r


def add_par(doc, align=None, space_before=0, space_after=4,
            first_line=None, left=None, hanging=None, keep_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    if left is not None:
        pf.left_indent = Inches(left)
    if first_line is not None:
        pf.first_line_indent = Inches(first_line)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)
    pf.widow_control = True
    if keep_next:
        pf.keep_with_next = True
    return p


def emit_subscripted(p, text, size=BODY_PT, italic=False):
    """Write text into p, turning _[x] into a real subscript run."""
    pos = 0
    for m in SUB.finditer(text):
        if m.start() > pos:
            style_run(p.add_run(text[pos:m.start()]), size, italic=italic)
        r = style_run(p.add_run(m.group(1)), size, italic=italic)
        r.font.subscript = True
        pos = m.end()
    if pos < len(text):
        style_run(p.add_run(text[pos:]), size, italic=italic)


def emit_marked(p, text, size=BODY_PT):
    """Write text into p, treating •...• spans as italic (used for journal titles)."""
    for i, chunk in enumerate(text.split("•")):
        if chunk:
            style_run(p.add_run(chunk), size, italic=bool(i % 2))


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_cell_text(cell, text, bold=False, size=CAPTION_PT):
    cell.text = ""
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    emit_subscripted(p, C.cite(text), size)
    for r in p.runs:
        r.bold = bold


# --------------------------------------------------------------------------- #
# document
# --------------------------------------------------------------------------- #

def build():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)

    # --- page setup: A4, 1in margins, ONE column ---------------------------- #
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1.0))
    cols = sec._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sec._sectPr.append(cols)
    cols.set(qn("w:num"), "1")
    cols.set(qn("w:space"), "360")

    # --- title block -------------------------------------------------------- #
    p = add_par(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    style_run(p.add_run(C.TITLE), 18)

    for name, sub in C.AUTHORS:
        p = add_par(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        style_run(p.add_run(name), BODY_PT)
        p = add_par(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        style_run(p.add_run(sub), BODY_PT, italic=True)

    for i, line in enumerate(C.AFFILIATION):
        p = add_par(doc, WD_ALIGN_PARAGRAPH.CENTER, space_after=(10 if i == len(C.AFFILIATION) - 1 else 0))
        style_run(p.add_run(line), BODY_PT)

    # --- abstract / index terms --------------------------------------------- #
    p = add_par(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    style_run(p.add_run("Abstract—"), SMALL_PT, bold=True, italic=True)
    style_run(p.add_run(C.cite(C.ABSTRACT)), SMALL_PT, bold=True, italic=True)

    p = add_par(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=10)
    style_run(p.add_run("Index Terms—"), SMALL_PT, bold=True, italic=True)
    style_run(p.add_run(C.INDEX_TERMS), SMALL_PT, bold=True, italic=True)

    # --- body --------------------------------------------------------------- #
    num_counter = 0
    for block in C.BODY:
        kind = block[0]

        if kind != "num":
            num_counter = 0

        if kind == "h1":
            p = add_par(doc, WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=5,
                        keep_next=True)
            style_run(p.add_run("%s  %s" % (block[1], block[2])), 12, bold=True)

        elif kind == "h2":
            p = add_par(doc, space_before=8, space_after=4, keep_next=True)
            style_run(p.add_run("%s  %s" % (block[1], block[2])), BODY_PT, italic=True)

        elif kind == "p":
            p = add_par(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=0.25, space_after=5)
            emit_subscripted(p, C.cite(block[1]))

        elif kind == "bullet":
            p = add_par(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, left=0.55, hanging=0.3, space_after=4)
            style_run(p.add_run("•\t"))
            emit_subscripted(p, C.cite(block[1]))

        elif kind == "num":
            num_counter += 1
            p = add_par(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, left=0.55, hanging=0.3, space_after=4)
            style_run(p.add_run("%d)\t" % num_counter))
            emit_subscripted(p, C.cite(block[1]))

        elif kind == "eq":
            p = add_par(doc, space_before=6, space_after=6, keep_next=True)
            pf = p.paragraph_format
            pf.tab_stops.add_tab_stop(Inches(MEASURE / 2.0), WD_TAB_ALIGNMENT.CENTER)
            pf.tab_stops.add_tab_stop(Inches(MEASURE), WD_TAB_ALIGNMENT.RIGHT)
            style_run(p.add_run("\t"))
            emit_subscripted(p, block[1], BODY_PT, italic=True)
            style_run(p.add_run("\t"))
            style_run(p.add_run(block[2]))

        elif kind == "where":
            p = add_par(doc, space_before=0, space_after=3)
            style_run(p.add_run("where"))
            for symbol, definition in block[1]:
                p = add_par(doc, WD_ALIGN_PARAGRAPH.JUSTIFY,
                            left=0.85, hanging=0.5, space_after=3)
                emit_subscripted(p, symbol, BODY_PT, italic=True)
                style_run(p.add_run("  —  "))
                emit_subscripted(p, C.cite(definition))

        elif kind == "fig":
            p = add_par(doc, WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=3,
                        keep_next=True)
            p.add_run().add_picture(os.path.join(HERE, block[1]),
                                    width=Inches(block[2]), height=Inches(block[3]))

        elif kind == "cap":
            p = add_par(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=10)
            style_run(p.add_run(block[1] + "  "), CAPTION_PT, bold=True)
            emit_subscripted(p, C.cite(block[2]), CAPTION_PT)

        elif kind == "tcap":
            p = add_par(doc, WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=4,
                        keep_next=True)
            style_run(p.add_run("%s %s" % (block[1], block[2])), CAPTION_PT)

        elif kind == "table":
            rows = block[1]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = True
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    set_cell_text(t.cell(ri, ci), cell, bold=(ri == 0))
            set_repeat_header(t.rows[0])
            add_par(doc, space_after=8)

        else:
            raise ValueError("unknown block kind %r" % kind)

    # --- references ---------------------------------------------------------- #
    p = add_par(doc, WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=6, keep_next=True)
    style_run(p.add_run("References"), 12, bold=True)

    for entry in C.references_sorted():
        p = add_par(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, left=0.5, hanging=0.5, space_after=4)
        emit_marked(p, entry, SMALL_PT)

    doc.save(OUT)
    print("wrote %s" % OUT)


if __name__ == "__main__":
    build()

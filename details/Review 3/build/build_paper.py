# -*- coding: utf-8 -*-
"""Build Review3_Research_Paper.docx (and .pdf) from paper_content.py.

The layout follows the Review 2 paper: US Letter, single column, a running header with the
department line and the page number, a centred title block, Roman-numbered sections,
lettered subsections, booktabs-style tables with captions above, figures with captions
below, numbered equations (native Word equations), boxed algorithms, IEEE references.

    python3 build_paper.py            (from details/Review 3/build)
"""

import html
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import omml  # noqa: E402
import paper_content as C  # noqa: E402

OUT_DIR = os.path.normpath(os.path.join(HERE, ".."))
OUT = os.path.join(OUT_DIR, "Review3_Research_Paper.docx")

FONT = "Times New Roman"
BODY = 10.0
SMALL = 9.0
TABLE = 8.5
REF = 8.0
PAGE_W, PAGE_H = 8.5, 11.0
MARGIN_LR = 0.68
TEXT_W = PAGE_W - 2 * MARGIN_LR

ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV", "XVI", "XVII", "XVIII", "XIX", "XX"]


# --------------------------------------------------------------------------- #
# numbering: labels and citations
# --------------------------------------------------------------------------- #

LABELS = {}
CITES = {}


def number_labels():
    counts = {"fig": 0, "tab": 0, "eq": 0, "alg": 0, "sec": 0}
    for blk in C.BODY:
        kind = blk[0]
        if kind == "h1":
            counts["sec"] += 1
            LABELS[blk[2]] = ROMAN[counts["sec"] - 1]
        elif kind in ("fig", "table", "eq", "alg"):
            key = {"fig": "fig", "table": "tab", "eq": "eq", "alg": "alg"}[kind]
            counts[key] += 1
            LABELS[blk[1]] = counts[key]


def all_texts():
    """Every string of the body in reading order (for numbering citations)."""
    for blk in C.BODY:
        kind = blk[0]
        if kind in ("p",):
            yield blk[1]
        elif kind in ("bullets", "enum"):
            yield from blk[1]
        elif kind == "fig":
            yield blk[4]
        elif kind == "table":
            yield blk[2]
            for row in blk[4]:
                yield from row
        elif kind == "alg":
            for _, t in blk[3]:
                yield t


CITE_RE = re.compile(r"\{([a-z]+[0-9]{4}(?:;[a-z]+[0-9]{4})*)\}")
REF_RE = re.compile(r"\[\[(fig|tab|eq|alg|sec):([a-z0-9_]+)\]\]")


REVIEW2_ORDER = ["mahmud2022", "gupta2017", "hasan2024", "rajab2021", "syed2022", "naim2016", "reyana2023",
                 "gupta2025", "pearl1988", "fenton2007", "weber2012", "aguilera2011", "pang2025"]


def number_citations():
    """Review 2's references keep their numbers [1]-[13]; new ones follow in order of first citation."""
    for i, k in enumerate(REVIEW2_ORDER, 1):
        CITES[k] = i
    n = len(REVIEW2_ORDER)
    for t in all_texts():
        for m in CITE_RE.finditer(t):
            for k in m.group(1).split(";"):
                if k not in C.REFERENCES:
                    raise KeyError("unknown citation %s" % k)
                if k not in CITES:
                    n += 1
                    CITES[k] = n
    unused = set(C.REFERENCES) - set(CITES)
    if unused:
        raise ValueError("references never cited: %s" % sorted(unused))


def expand(text):
    def cite(m):
        return ", ".join("[%d]" % CITES[k] for k in m.group(1).split(";"))

    def ref(m):
        kind, name = m.group(1), m.group(2)
        lab = LABELS["%s:%s" % (kind, name)]
        return {"fig": "Fig. %s", "tab": "Table %s", "eq": "(%s)", "alg": "Algorithm %s",
                "sec": "Section %s"}[kind] % (lab if kind != "tab" else ROMAN[lab - 1])

    text = CITE_RE.sub(cite, text)
    text = REF_RE.sub(ref, text)
    return text


# --------------------------------------------------------------------------- #
# run helpers
# --------------------------------------------------------------------------- #

def set_font(r, size, bold=None, italic=None, caps=False, smallcaps=False, color=None):
    r.font.name = FONT
    r.font.size = Pt(size)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    if caps:
        r.font.all_caps = True
    if smallcaps:
        r.font.small_caps = True
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    return r


TAG_RE = re.compile(r"(</?(?:i|b|sub|sup)>)")


def emit(p, text, size=BODY, bold=False, italic=False):
    """Write marked-up text (<i>, <b>, <sub>, <sup>) into paragraph p."""
    text = expand(text)
    st = {"i": italic, "b": bold, "sub": False, "sup": False}
    for tok in TAG_RE.split(text):
        if not tok:
            continue
        m = re.fullmatch(r"<(/?)(i|b|sub|sup)>", tok)
        if m:
            st[m.group(2)] = not m.group(1)
            continue
        r = p.add_run(html.unescape(tok))
        set_font(r, size, bold=st["b"], italic=st["i"])
        if st["sub"]:
            r.font.subscript = True
        if st["sup"]:
            r.font.superscript = True


def para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, first=None, left=None,
         hanging=None, keep=False, line=None, right=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line if line is not None else 1.0
    if left is not None:
        pf.left_indent = Inches(left)
    if right is not None:
        pf.right_indent = Inches(right)
    if first is not None:
        pf.first_line_indent = Inches(first)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)
    pf.widow_control = True
    if keep:
        pf.keep_with_next = True
    return p


def border(p, sides, sz=8, color="000000"):
    pPr = p._p.get_or_add_pPr()
    pb = pPr.find(qn("w:pBdr"))
    if pb is None:
        pb = OxmlElement("w:pBdr")
        after = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:contextualSpacing",
                 "w:jc", "w:textDirection", "w:textAlignment", "w:outlineLvl", "w:rPr")
        nxt = next((c for c in pPr if c.tag in [qn(a) for a in after]), None)
        if nxt is not None:
            nxt.addprevious(pb)
        else:
            pPr.append(pb)
    for side in sides:
        el = OxmlElement("w:%s" % side)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "1")
        el.set(qn("w:color"), color)
        pb.append(el)


def cell_borders(cell, **sides):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:%s" % side)
        sz = sides.get(side)
        if sz:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        tb.append(el)
    tcPr.append(tb)


def cell_margins(table, top=30, bottom=30, left=70, right=70):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:%s" % side)
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    lay = tblPr.find(qn("w:tblLayout"))     # python-docx adds it for autofit = False
    if lay is None:
        lay = OxmlElement("w:tblLayout")
        tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    lay.addnext(m)


def no_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)
    if header:
        h = OxmlElement("w:tblHeader")
        h.set(qn("w:val"), "true")
        trPr.append(h)


def field(p, instr, size):
    """Insert a simple field (e.g. PAGE)."""
    r = p.add_run()
    set_font(r, size)
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r._r.append(fc)
    r2 = p.add_run()
    set_font(r2, size)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = p.add_run()
    set_font(r3, size)
    fc2 = OxmlElement("w:fldChar")
    fc2.set(qn("w:fldCharType"), "separate")
    r3._r.append(fc2)
    r4 = p.add_run("1")
    set_font(r4, size)
    r5 = p.add_run()
    fc3 = OxmlElement("w:fldChar")
    fc3.set(qn("w:fldCharType"), "end")
    r5._r.append(fc3)


# --------------------------------------------------------------------------- #
# document
# --------------------------------------------------------------------------- #

def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-GB")
    rpr.append(lang)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(PAGE_W), Inches(PAGE_H)
    sec.left_margin = sec.right_margin = Inches(MARGIN_LR)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.75)
    sec.header_distance = Inches(0.36)
    sec.footer_distance = Inches(0.4)

    hp = sec.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run(C.HEADER)
    set_font(r, 7)
    r = hp.add_run("\t")
    set_font(r, 7)
    field(hp, "PAGE", 7)

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")

    props = doc.core_properties
    props.title = C.TITLE
    props.author = "Rahul Akshath Rajendran, Mohit R"
    props.subject = "Project Review 3 research paper, Department of Information Technology, SSN College of Engineering"
    props.keywords = C.INDEX_TERMS


def title_block(doc):
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=14, line=1.0)
    set_font(p.add_run(C.TITLE), 24)
    for name, lines in C.AUTHORS:
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=1)
        set_font(p.add_run(name), 11)
        for i, ln in enumerate(lines):
            p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=(9 if i == len(lines) - 1 else 1))
            set_font(p.add_run(ln), 8)
    for i, ln in enumerate(C.AFFILIATION):
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=(16 if i == len(C.AFFILIATION) - 1 else 0))
        set_font(p.add_run(ln), 9)

    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=4, keep=True)
    set_font(p.add_run("Abstract"), SMALL)
    p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, after=10, left=0.45, right=0.45, first=0.17)
    emit(p, C.ABSTRACT, SMALL)
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=4, keep=True)
    set_font(p.add_run("Index Terms"), SMALL)
    p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, left=0.45, right=0.45, first=0.17)
    emit(p, C.INDEX_TERMS, SMALL)


def dropcap_paragraph(doc, text):
    """First paragraph of the introduction with a two-line drop cap, as in IEEEtran."""
    text = expand(text)
    first, rest = text[0], text[1:]
    m = re.match(r"(\S+)(.*)", rest, re.S)
    word_rest, tail = m.group(1), m.group(2)
    dc = para(doc, WD_ALIGN_PARAGRAPH.LEFT, line=1.0)
    pPr = dc._p.get_or_add_pPr()
    fp = OxmlElement("w:framePr")
    fp.set(qn("w:dropCap"), "drop")
    fp.set(qn("w:lines"), "2")
    fp.set(qn("w:wrap"), "around")
    fp.set(qn("w:vAnchor"), "text")
    fp.set(qn("w:hAnchor"), "text")
    pPr.insert(0, fp)
    r = dc.add_run(first)
    set_font(r, 23.5)
    rPr = r._element.get_or_add_rPr()
    pos = OxmlElement("w:position")
    pos.set(qn("w:val"), "-3")
    rPr.append(pos)
    p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    set_font(p.add_run(word_rest.upper()), BODY)
    emit(p, tail)


def build():
    number_labels()
    number_citations()
    doc = Document()
    setup(doc)
    title_block(doc)

    sub = 0
    for blk in C.BODY:
        kind = blk[0]
        if kind == "h1":
            sub = 0
            p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6, keep=True)
            r = p.add_run("%s. %s" % (LABELS[blk[2]], blk[1]))
            set_font(r, BODY, smallcaps=True)
        elif kind == "h2":
            p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=8, after=3, keep=True)
            r = p.add_run("%s. %s" % ("ABCDEFGHIJKLMNOP"[sub], blk[1]))
            set_font(r, BODY, italic=True)
            sub += 1
        elif kind == "p":
            if len(blk) > 2 and blk[2] == "dropcap":
                dropcap_paragraph(doc, blk[1])
            else:
                p = para(doc, first=0.17)
                emit(p, blk[1])
        elif kind == "bullets":
            for item in blk[1]:
                p = para(doc, left=0.42, hanging=0.17, before=1, after=1)
                set_font(p.add_run("•\t"), BODY)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(0.42))
                emit(p, item)
        elif kind == "enum":
            for i, item in enumerate(blk[1], 1):
                p = para(doc, left=0.42, hanging=0.25, before=1, after=1)
                set_font(p.add_run("%d)\t" % i), BODY)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(0.42))
                emit(p, item)
        elif kind == "eq":
            p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=5, after=5)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
            set_font(p.add_run("\t"), BODY)
            p._p.append(parse_xml(omml.omath(blk[2])))
            set_font(p.add_run("\t(%d)" % LABELS[blk[1]]), BODY)
        elif kind == "fig":
            _, lab, path, width, cap = blk
            p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, keep=True)
            p.add_run().add_picture(os.path.join(HERE, path), width=Inches(width))
            p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, after=10)
            set_font(p.add_run("Fig. %d: " % LABELS[lab]), SMALL)
            emit(p, cap, SMALL)
        elif kind == "table":
            render_table(doc, blk)
        elif kind == "alg":
            render_alg(doc, blk)
        else:
            raise ValueError(kind)

    # references
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6, keep=True)
    set_font(p.add_run("References"), BODY, smallcaps=True)
    for key, n in sorted(CITES.items(), key=lambda kv: kv[1]):
        p = para(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, left=0.32, hanging=0.32, after=2)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.32))
        set_font(p.add_run("[%d]\t" % n), REF)
        emit(p, C.REFERENCES[key], REF)

    doc.save(OUT)
    print("wrote", OUT)


def render_table(doc, blk):
    _, lab, title, header, rows, widths = blk[:6]
    compact = len(blk) > 6 and blk[6] == "compact"
    size = TABLE if compact else 9.0
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=10, after=5, keep=True)
    set_font(p.add_run("TABLE %s: " % ROMAN[LABELS[lab] - 1]), SMALL)
    emit(p, title, SMALL)
    scale = min(1.0, TEXT_W / sum(widths))
    widths = [w * scale for w in widths]
    t = doc.add_table(rows=len(rows) + 1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t, top=25, bottom=25, left=60, right=60)
    for ci, w in enumerate(widths):
        t.columns[ci].width = Inches(w)
    grid = t._tbl.tblGrid
    for ci, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[ci] * 1440)))
    numeric = [all(re.fullmatch(r"[<>&;a-z\s]*[−+\-]?[\d.,]+[%]?(?:\s*\[.*\])?|n/a|—|[\d.–]+", re.sub(r"</?b>", "", r[c]).strip() or "x")
                   for r in rows) for c in range(len(header))]
    for ri in range(len(rows) + 1):
        row = t.rows[ri]
        no_split(row, header=(ri == 0))
        for ci in range(len(header)):
            cell = row.cells[ci]
            cell.width = Inches(widths[ci])
            text = header[ci] if ri == 0 else rows[ri - 1][ci]
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.line_spacing = 1.0
            if ci > 0 and numeric[ci] and ri > 0:
                cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif ri == 0 and ci > 0 and numeric[ci]:
                cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            emit(cp, text, size)
            if ri < len(rows) and len(rows) <= 14:
                cp.paragraph_format.keep_with_next = True
            if ri == 0:
                for r in cp.runs:
                    r.bold = False
            top = 12 if ri == 0 else None
            bottom = 6 if ri == 0 else (12 if ri == len(rows) else None)
            cell_borders(cell, top=top, bottom=bottom)
    para(doc, after=6)


def render_alg(doc, blk):
    _, lab, title, lines = blk
    p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=8, after=0, keep=True)
    border(p, ["top", "bottom"], sz=8)
    set_font(p.add_run("Algorithm %d: " % LABELS[lab]), SMALL, bold=True)
    emit(p, title, SMALL)
    for i, (ind, text) in enumerate(lines, 1):
        last = i == len(lines)
        p = para(doc, WD_ALIGN_PARAGRAPH.LEFT, after=(0 if not last else 8), left=0.3 + 0.22 * ind,
                 hanging=0.3 + 0.22 * ind if ind == 0 else None, keep=not last)
        if ind:
            p.paragraph_format.first_line_indent = Inches(-(0.3 + 0.22 * ind))
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3 + 0.22 * ind))
        lead = "%d:\t" % i + ("\t" if ind else "")
        set_font(p.add_run(lead), SMALL)
        emit(p, text, SMALL)
        if last:
            border(p, ["bottom"], sz=8)


def to_pdf():
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT_DIR, OUT],
                   check=True, capture_output=True, timeout=300)
    print("wrote", os.path.splitext(OUT)[0] + ".pdf")


if __name__ == "__main__":
    build()
    if "--no-pdf" not in sys.argv:
        to_pdf()
